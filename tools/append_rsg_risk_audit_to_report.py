"""Append source-backed RSG objective and frozen-expert audits to the report."""
import json
from pathlib import Path
from statistics import mean

from docx import Document
from append_phr_candidate_to_official_report import add_body, add_heading, set_run_font
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
TITLE = '附录 J 门控训练目标的理论边界与固定专家验证'


def equation(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math = OxmlElement('m:oMath')
    run = OxmlElement('m:r')
    value = OxmlElement('m:t')
    value.text = text
    run.append(value)
    math.append(run)
    paragraph._p.append(math)


def main():
    report = next((ROOT/'结题').glob('*（正式版）.docx'))
    document = Document(report)
    if any(p.text == TITLE for p in document.paragraphs):
        print('already_present')
        return
    counter = json.loads((ROOT/'outputs/theory/rsg_population_counterexample.json').read_text())
    pilot = json.loads((ROOT/'outputs/training/frozen_expert_gate_pilot/summary.json').read_text())
    variational = json.loads((ROOT/'outputs/theory/rsg_variational_target.json').read_text())
    if not (counter['all_checks_passed'] and variational['all_checks_passed'] and pilot['frozen_expert_parameters_and_buffers_unchanged']):
        raise ValueError('Required verification did not pass')
    image_count = len(document.inline_shapes)
    table_count = len(document.tables)
    document.add_page_break()
    add_heading(document, TITLE, 1)
    add_body(document, '本附录补充 RSG-HRGV 门控目标的解析依据、总体风险边界及固定专家对照。已有路由后悔下降结果不能直接解释为门控的独立性能收益，更不能推出整体分类指标必然提高。以下内容用于限定网络创新的可辩护范围，并据此确定后续改进问题。')
    add_heading(document, 'J.1 软门控目标的优化解释', 2)
    add_body(document, '固定单张图像和真实角色，两专家对数损失记为 ld、lm，温度 T>0。考虑专家期望损失与负熵正则之和：')
    equation(document, 'Jₜ(t) = t l_d + (1−t) l_m + T[t log t + (1−t) log(1−t)]')
    add_body(document, '内部二阶导数为 T/[t(1−t)]>0，由一阶导数为零和端点导数极限，可得唯一最优解 t*=sigmoid((lm−ld)/T)，与已实现软目标相同。该解释是经典熵正则化选择原理的两专家实例化，不作为独立首创定理。原理参考 Niculae 与 Blondel，A Regularized Framework for Sparse and Structured Neural Attention，NeurIPS 2017。')
    equation(document, 'Jₜ(g) − Jₜ(t*) = T KL(Ber(g) ∥ Ber(t*))')
    equation(document, 'BCE(g,t*) − H(t*) = KL(Ber(t*) ∥ Ber(g))')
    add_body(document, '两式的 KL 方向不同。因此 BCE 是对软 oracle 的拟合，不能被写成直接优化熵正则目标的等价实现。四组温度各 3,000 个合成样本的数值核验均通过，其中包含当前温度 0.2。数值核验支持公式实现一致性，不替代解析证明或分类性能实验。')
    add_heading(document, 'J.2 总体目标非等价性的反例', 2)
    add_body(document, '对同一输入 x，推理门控不能使用未知真实角色。记真实条件分布为 π，标签相关的软目标为 ty、权重为 wy。在条件平均权重为正时，总体归一化加权 BCE 的最优门控为：')
    equation(document, 'g*BCE = E[wY tY | x] / E[wY | x]')
    add_body(document, '而融合风险为 R(g)=−Σy πy log[by+g(ay−by)]，一般不在上述门控处最小。取 π=(0.4,0.4,0.1,0.1)，a=(0.72,0.08,0.1,0.1)，b=(0.48,0.32,0.1,0.1)。各概率严格为正，使用当前软目标及差距权重。')
    equation(document, "R′(g) = 0.096[1/(0.32−0.24g) − 1/(0.48+0.24g)] > 0")
    add_body(document, f"对 g∈[0,1]，上式严格为正，故融合 NLL 唯一最优门控为 0。实际代码计算的 BCE 最优门控为 {counter['weighted_bce_optimal_gate']:.6f}，额外产生 {counter['excess_nll_nats']:.6f} nat 的条件 NLL。这个有限支持反例否定无条件的 NLL 一致性主张，但未模拟完整联合训练网络，也不证明当前矿物数据上的真实失效原因。")
    add_heading(document, 'J.3 固定专家的验证集对照', 2)
    add_body(document, '固定一个未采用后悔监督的 HRGV 专家检查点，在训练集上重新训练门控。比较等权融合、直接融合 NLL 和后悔 BCE 三种策略。两种学习门控均使用三组配对初始化、相同 30 轮预算，按验证集融合 NLL 选轮次。专家全部参数和缓冲区的逐项精确相等检查通过。')
    table = document.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, ('策略', '融合 NLL / nat', '路由后悔 / nat', 'Accuracy')):
        cell.text = text
    records = [('等权融合', pilot['equal'])]
    for prefix, label in (('fusion_nll', '融合 NLL 门控'), ('regret_bce', '后悔 BCE 门控')):
        runs = [v for k, v in pilot['runs'].items() if k.startswith(prefix)]
        if len(runs) != 3:
            raise ValueError('Expected three paired initializations')
        records.append((label, {key: mean(r[key] for r in runs) for key in ('fusion_nll_nats', 'routing_regret_nats', 'accuracy')}))
    for label, values in records:
        for cell, text in zip(table.add_row().cells, (label, f"{values['fusion_nll_nats']:.6f}", f"{values['routing_regret_nats']:.6f}", f"{values['accuracy']:.2%}")):
            cell.text = text
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, '宋体', 10)
    add_body(document, '后悔 BCE 相对等权融合的验证 NLL 仅降低约 0.001381 nat，Accuracy 下降约 0.1817 个百分点。最佳轮次集中在前两轮。三组种子只改变门控初始化，不是三个独立专家。源专家已拟合训练样本并曾用当前验证集选模型，本实验再用同一验证集选门控轮次，故结果仅为探索性机制诊断，不能作独立泛化或显著性结论。')
    add_heading(document, 'J.4 网络改进需要解决的问题', 2)
    add_body(document, '后续应分别检验两项问题：其一，训练内专家预测能否代表未见样本上的门控监督；其二，门控训练目标与最终融合风险是否一致。交叉拟合可用于研究第一项，但不能单独消除第二项的总体目标非等价性。新的网络方法应在明确的最终风险目标下提出，再用固定专家对照和独立数据检验，当前不宣称该改进已完成。')
    add_body(document, '复现依据：verify_rsg_variational_target.py、verify_rsg_population_counterexample.py、audit_frozen_expert_routing.py 与 run_frozen_expert_gate_pilot.py；数值结果位于 outputs/theory，门控实验协议、全部轮次和验证预测位于 outputs/training/frozen_expert_gate_pilot。本文对数损失统一使用 nat；将损失乘以100展示也不能解释为概率百分点。')
    temporary = report.with_suffix('.risk_audit.tmp.docx')
    document.save(temporary)
    checked = Document(temporary)
    assert len(checked.inline_shapes) == image_count
    assert len(checked.tables) == table_count+1
    assert sum(p.text == TITLE for p in checked.paragraphs) == 1
    temporary.replace(report)
    print(f'updated paragraphs={len(checked.paragraphs)} tables={len(checked.tables)} images={image_count}')


if __name__ == '__main__':
    main()
