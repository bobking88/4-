"""Append the Fold 0 capacity diagnosis without rewriting the official report."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, RGBColor

from append_abmp_to_official_report import _add_body, _add_heading, _equation, _table
from append_oos_rsg_to_official_report import _set_run_font


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "结题" / "基于深度学习的钒钛矿相关矿物图像识别方法研究_技术报告（正式版）.docx"
TITLE = "附录 N 候选空间容量与校验约束的理论分析"


def _hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def append_capacity_appendix(report, summary, properties, figure_dir, *, expected_sha256):
    if _hash(report) != expected_sha256:
        raise ValueError("Report hash changed; refusing to edit another version.")
    if summary.get("protocol") != "abmp_candidate_capacity_v1" or summary.get("status") != "diagnostic_complete_no_promotion" or summary.get("outer_images_loaded") is not False or summary.get("fold_1_2_accessed") is not False:
        raise ValueError("Only the registered inner development evidence may be appended.")
    if properties.get("purpose") != "structural_formula_check_not_classifier_performance" or any(properties.get(key) != 0 for key in ("nested_segment_violations", "target_set_promotions", "scale_margin_violations")):
        raise ValueError("Geometry checks have missing or failed invariants.")
    # This appendix explains this specific development failure, not future audits.
    for name, expected in (("gate_stop_projector_fit", (35, 35, 38, 38, 2, 3)), ("projector_stop", (43, 43, 43, 45, 0, 4))):
        subset = summary["subsets"][name]
        count = subset["class_counts"]["0"]
        observed = (round(subset["metrics"]["fixed_verified"]["target_recall"]*count),
                    *[subset["segment_capacity"][key]["true_target_feasible"] for key in ("restricted_verified", "full_verified", "full_pre")],
                    subset["verification_effects"]["fixed"]["true_target_removed"], subset["verification_effects"]["fixed"]["false_target_removed"])
        residual = subset["verifier_commutation_max_abs_residual"]
        if count != 66 or subset["sample_count"] != 340 or observed != expected or subset["nested_segment_violation_count"] != 0 or not math.isfinite(residual) or residual > 1e-12 or any(item["target_promotions"] != 0 for item in subset["verification_effects"].values()):
            raise ValueError("Audit observations changed; revise the narrative before appending.")
    document = Document(report)
    if any(p.text == TITLE for p in document.paragraphs):
        return False
    equation_names = ("scale", "commutation", "monotonicity", "interval", "capacity", "protected_scale")
    formulas = {name: figure_dir / "capacity_formulas" / f"capacity_{name}.png" for name in equation_names}
    figure = figure_dir / "fig_abmp_candidate_capacity.png"
    for path in [figure, *formulas.values()]:
        if not path.is_file():
            raise FileNotFoundError(path)
    fit = summary["subsets"]["gate_stop_projector_fit"]
    stop = summary["subsets"]["projector_stop"]

    document.add_page_break()
    _add_heading(document, TITLE, 1)
    _add_body(document, "本附录解释 ABMP-RSG-Net v2 开发门未通过的结构原因，并给出下一轮网络改进的可检验依据。数学性质、标签辅助容量诊断与模型性能结论分别陈述。当前尚未证明 v2 优于既有基线，也未启动新结构的独立确证。")
    _add_heading(document, "N.1 开发范围和候选容量证据", 2)
    _add_body(document, f"仅读取 Fold 0 的策略拟合集 gate_stop_projector_fit（{fit['sample_count']} 张）与停止集 projector_stop（{stop['sample_count']} 张），每个子集含 66 张目标代理样本。专家和候选门权重保持冻结；未读取 outer_eval 清单或图像，未访问 Fold 1/2。图像 ID、重复组及图片编号的跨子集交集均为 0。")
    rows = []
    for name, subset in (("策略拟合集", fit), ("策略停止集", stop)):
        observed = round(subset["metrics"]["fixed_verified"]["target_recall"]*subset["class_counts"]["0"])
        rows.append([name, observed, *[subset["segment_capacity"][key]["true_target_feasible"] for key in ("restricted_verified", "full_verified", "full_pre")]])
    _table(document, ["子集", "固定融合\n实际命中", "受限区间\n可达目标", "完整校验区间\n可达目标", "校验前完整区间\n可达目标"], rows, [3.0, 2.8, 3.2, 3.5, 3.5])
    table = document.tables[-1]
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn(f"w:{key}"), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(str(figure), width=Cm(16))
    caption = document.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run("图 N-1 开发子集的目标容量与固定融合校验损失")
    _set_run_font(run, "宋体", 10)
    run.font.color.rgb = RGBColor(0, 0, 0)
    _add_body(document, "可达数量由真实标签辅助、逐图选择最有利融合比例计算，仅是有限样本上界，不是可部署网络的正确率。校验后停止集的完整区间仍只有 43 张目标可达；校验前为 45 张。扩大路由范围在该子集不能新增目标命中，而校验尺度对 2 张潜在恢复样本有影响。")

    document.add_page_break()
    _add_heading(document, "N.2 校验尺度与目标集合的解析性质", 2)
    _add_body(document, "设 d、m 为直接角色后验和种类映射后验，目标类别索引为 T=0。两路专家对同一图像使用相同的残差校验尺度 s；s 由两类校验器的矛盾程度经指数变换得到。s 严格为正且不大于 1，后验均经过归一化。")
    _equation(document, formulas["scale"], "N-1")
    _add_body(document, "性质 1：共享校验不形成新的专家混合族，只改变混合权重的参数化。对 g∈[0,1]，将未归一化分子展开为 gDsd+(1−g)Dsm，再分别代入 Z(d)、Z(m)，即可得到式 N-2。Z 为正，因此 g′ 随 g 连续严格递增并覆盖 [0,1]。")
    _equation(document, formulas["commutation"], "N-2")
    _add_body(document, "需要注意，校验后的等效权重 g′ 通常不等于原门控输出 g。若忽略两路归一化因子，会错误解释校验后融合比例。正尺度映射可逆，因而不能将其一般性地描述为不可逆的信息丢失，也不能未经证明宣称其必然收缩所有分歧或 JS 散度。")
    _add_body(document, "性质 2：单向校验只能保持或移除目标预测，不能把原非目标预测提升为目标。若校验后目标获胜，则 spT 不小于任何 pk；由于 s≤1，原 pT 也不小于 pk。采用目标索引优先的固定 argmax 平局规则时，目标预测集合具有包含关系。")
    _equation(document, formulas["monotonicity"], "N-3")
    _add_body(document, f"这意味着单向校验后的目标召回及各非目标误入率都不会增加，但总体准确率和 Precision 不具有同样的单调保证。固定融合在拟合集移除 {fit['verification_effects']['fixed']['true_target_removed']} 张真实目标、{fit['verification_effects']['fixed']['false_target_removed']} 张假目标；停止集对应为 {stop['verification_effects']['fixed']['true_target_removed']} 和 {stop['verification_effects']['fixed']['false_target_removed']} 张。降低误入率必须与真实目标损失一并报告。")

    document.add_page_break()
    _add_heading(document, "N.3 精确可行区间与路由能力界", 2)
    _add_body(document, "不能只用两端 argmax 的并集判断容量。两个非目标端点若分别偏向不同非目标类别，中间混合仍可能使目标获胜。例如 d=(0.34,0.50,0.08,0.08)、m=(0.34,0.08,0.50,0.08) 的两端都不判为目标，而 g 在约 0.381 至 0.619 之间时目标可以获胜。")
    _equation(document, formulas["interval"], "N-4")
    _add_body(document, "性质 3：目标间隔条件是关于 g 的线性不等式。对每个非目标类别，若 bk>0，则增加下界 (η−ak)/bk；若 bk<0，则降低上界；若 bk=0 且 ak<η，区间为空。与 [0,1] 相交后得到精确闭区间；下界高于上界即为当前两路后验无恢复能力的证书。实现以 float64 计算并检查常数方向。")
    _add_body(document, "q0 为共享校验后的固定平均融合，qφ 为共享校验后的 OOS 门控融合。由性质 1，二者均位于完整校验专家线段上；二者之间的受限线段也必然属于完整线段。定义 Cη 为存在至少 η 目标间隔的图像集合，则得到式 N-5。")
    _equation(document, formulas["capacity"], "N-5")
    _add_body(document, "RecT,η 表示预测目标且类别间隔至少为 η 的真实目标比例。η=0 并采用目标优先平局规则时，对应普通目标召回；本次数值审计使用 η=10⁻⁹，排除极小间隔的数值歧义。容量上界依赖已冻结的两路后验，不能解释为任意神经网络的泛化上界。")
    _add_body(document, "受限校验区间在两个开发子集均没有新增目标恢复空间。完整校验区间在拟合集另有 3 张可达目标，但停止集没有；这不足以支持将双向全区间路由单独作为改进方案。即使校验前放宽完整区间，停止集仍有 21/66 张目标无法恢复，说明部分问题来自专家表示而非门控范围。")
    _add_body(document, f"真实数据的等效权重重放最大绝对残差为 {max(fit['verifier_commutation_max_abs_residual'], stop['verifier_commutation_max_abs_residual']):.2e}，嵌套区间违反数为 0；另以随机种子 {properties['seed']} 验证 {properties['sample_count']:,} 对 float64 后验，目标集合、区间嵌套及校验间隔约束违反数均为 0。随机数值检查补充代数证明，但不替代分类实验。")

    document.add_page_break()
    _add_heading(document, "N.4 校验前间隔保护及后续网络研究", 2)
    _add_body(document, "现有 v2 在完整校验之后保护 q0 的锚点，不能直接恢复已被校验抑制的目标。下一阶段可检验将校验可信度与校验前间隔约束耦合。给定候选后验 p，若希望归一化校验后 pT−pk 至少为 δ，整理不等式可得最小尺度条件。")
    _equation(document, formulas["protected_scale"], "N-6")
    _add_body(document, "该条件要求 pT>0、0<δ<1；若 smin>1，则当前候选连不施加校验也不能达到指定间隔。若 smin≤1，可将学习得到的尺度投影至 [smin,1]。这是保持指定预测间隔的解析条件，不保证锚点标签为真；假目标锚点也会被保持。因此不能对所有图像无条件启用保护，更不能把样本级保证等同于真实矿物回收率保证。")
    _add_body(document, "后续实验先固定专家并比较校验关闭、固定强度、开发集标定强度和学习可信度四种对照；确认收益确实来自逐图可信度而非简单减弱校验。只有该对照通过，才结合完整路由与解析可行域开展网络实验，并将预算头、间隔保护和可信度门逐项消融。若专家容量仍不足，则需要学习互补的目标—难负样本残差表示，而不是再叠加仅改变后验混合比例的门。")
    _add_body(document, "方案进入独立确证前，应在 Fold 0 内冻结损失、超参数、锚点选择及失败判据；不得以 Fold 1/2 的结果反复改结构。评价同时包含 Macro F1、NLL、目标召回、两类难负样本误入率、真实及假锚点保持情况、约束激活率和参数量，并与同容量、同预算对照比较。当前方法尚无新模型性能结果。")
    _add_body(document, "理论贡献的定位是将候选可行域、校验的单向性和约束位置统一用于矿物网络的可证伪设计，而非宣称首次提出凸投影、层级学习或代价敏感分类。2026 年 ICML 的 Design Linear Constrained Neural Layers with Implicit Convex Optimization 已研究线性约束神经层；2026 年 CVPR 的 Every Error has Its Magnitude 已研究非对称错误严重程度和层级一致性。本次核验为官方论文摘要级查新，尚未完成与具体实现的全文对照。")
    _add_body(document, "参考入口：Yan 等，ICML 2026，https://proceedings.mlr.press/v306/yan26m.html；Hong 等，CVPR 2026，https://openaccess.thecvf.com/content/CVPR2026/html/Hong_Every_Error_has_Its_Magnitude_Asymmetric_Mistake_Severity_Training_for_CVPR_2026_paper.html。复现证据包括 audit_summary.json、逐图后验、routing_evidence.csv、geometry_properties.json 及图 N-1 的源数据。")

    # Save only after checking that no other writer changed the user's report.
    if _hash(report) != expected_sha256:
        raise ValueError("Report hash changed during construction; refusing to overwrite.")
    with tempfile.NamedTemporaryFile(dir=report.parent, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        document.save(temporary)
        if _hash(report) != expected_sha256:
            raise ValueError("Report hash changed before replace; refusing to overwrite.")
        temporary.replace(report)
    finally:
        temporary.unlink(missing_ok=True)
    return True


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", type=Path, default=REPORT)
    parser.add_argument("--audit-dir", type=Path, default=ROOT / "outputs/theory/abmp_candidate_capacity_v1")
    parser.add_argument("--figure-dir", type=Path, default=ROOT / "outputs/paper_figures_v5")
    parser.add_argument("--expected-sha256", required=True)
    args = parser.parse_args()
    summary = json.loads((args.audit_dir / "audit_summary.json").read_text(encoding="utf-8"))
    properties = json.loads((args.audit_dir / "geometry_properties.json").read_text(encoding="utf-8"))
    changed = append_capacity_appendix(args.report, summary, properties, args.figure_dir, expected_sha256=args.expected_sha256)
    print(json.dumps({"changed": changed, "sha256": _hash(args.report)}))


if __name__ == "__main__":
    main()
