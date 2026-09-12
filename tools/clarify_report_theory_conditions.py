"""Clarify theorem assumptions without changing report formulas or figures."""
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]


def revise(report):
    document = Document(report)
    replacements = {
        "给定校准后验 q、错误代价矩阵": "给定真实条件后验 q、已知错误代价矩阵 C，且拒识后的总损失为常数 crej 时，逐图比较所有直接动作的最小条件风险与 crej，可得贝叶斯最优选择性决策。使用网络估计后验时，该规则仅最小化估计后验下的风险；概率校准本身不足以保证真实条件风险最优。若拒识意味着送检，其风险还应包含检测后的残余决策损失。本研究仅验证公开标本图像上的选择性识别，不声称真实 XRF 送检成本已得到验证。",
        "若构造 g* 与 w 的专家后验使用停止梯度": "设门控输入 z 包含共享特征、两专家熵与专家分歧。只有对完整输入 z、软目标 g* 和权重 w 同时停止梯度，且门控参数与专家参数不共享时，后悔监督关于共享主干和专家参数的梯度才严格为零，门控自身参数梯度仍可非零。仅停止目标与权重的梯度不足以切断经门控输入反传的路径。实现中的 detach_gate_features=True 切断该路径。结论仅限 Lreg 分支，不中断角色、种类、验证器和对比损失对共享主干的训练。",
    }
    changed = 0
    for old, new in replacements.items():
        matches = [p for p in document.paragraphs if p.text.startswith(old)]
        if not matches and any(p.text == new for p in document.paragraphs):
            continue
        if len(matches) != 1:
            raise ValueError(f"Expected exactly one paragraph: {old}")
        paragraph = matches[0]
        if paragraph._p.xpath('.//w:drawing | .//m:oMath'):
            raise ValueError("Refusing to replace a figure or equation paragraph")
        paragraph.runs[0].text = new
        for run in paragraph.runs[1:]:
            run.text = ""
        changed += 1
    if changed:
        temporary = report.with_suffix('.theory_conditions.tmp.docx')
        document.save(temporary)
        reread = Document(temporary)
        assert len(reread.inline_shapes) == len(document.inline_shapes)
        assert len(reread.tables) == len(document.tables)
        temporary.replace(report)
    return changed


if __name__ == '__main__':
    report = next((ROOT / '结题').glob('*（正式版）.docx'))
    print(f'changed={revise(report)}')
