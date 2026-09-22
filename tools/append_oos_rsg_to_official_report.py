"""Append the OOS-RSG-HRGV candidate-network theory and pilot evidence."""
from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "结题" / "基于深度学习的钒钛矿相关矿物图像识别方法研究_技术报告（正式版）.docx"
FIGURE = ROOT / "outputs" / "paper_figures_v3" / "fig_oos_rsg_training_architecture.png"
ANALYSIS = ROOT / "outputs" / "training" / "seen_unseen_gate_study_v1" / "analysis.json"
DEGENERACY = ROOT / "outputs" / "theory" / "oos_gate_degeneracy.json"
BACKUP = ROOT / "结题" / "历史版本" / "基于深度学习的钒钛矿相关矿物图像识别方法研究_技术报告（正式版_20260922_追加OOS-RSG前）.docx"
TITLE = "附录 K 样本外风险监督路由网络及其理论性质"


def _set_run_font(run, name: str, size: float = 11.0, bold: bool = False, color: str | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for field in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{field}"), name)


def _add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(8 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(4)
    _set_run_font(paragraph.add_run(text), "黑体", 14 if level == 1 else 12, True)


def _add_body(document: Document, text: str, first_indent: bool = True) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.30
    if first_indent:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    _set_run_font(paragraph.add_run(text), "宋体", 11)


def _add_equation(document: Document, text: str, number: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    _set_run_font(run, "Cambria Math", 11)
    paragraph.add_run("    ")
    label = paragraph.add_run(f"({number})")
    _set_run_font(label, "Times New Roman", 10.5)


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_cell(cell, text: str, bold: bool = False, size: float = 9.0, shade: str | None = None) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    _set_run_font(paragraph.add_run(text), "宋体", size, bold)
    if shade:
        _shade(cell, shade)


def _keep_phr_result_header_with_data(document: Document) -> None:
    """Keep the existing PHR screening header with its first result row."""
    for table in document.tables:
        if len(table.rows) != 3 or len(table.columns) != 6:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header[:2] != ["配置", "Macro F1"] or "Ti" not in header[3]:
            continue
        _repeat_header(table.rows[0])
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
        for row in table.rows[:2]:
            properties = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            properties.append(cant_split)
        return


def _aggregate_map(analysis: dict) -> dict[tuple[str, str], dict]:
    return {(row["source"], row["objective"]): row for row in analysis["aggregates"]}


def _metric(row: dict, name: str, percent: bool = False) -> str:
    mean_value = row[f"mean_{name}"]
    sd_value = row[f"sample_sd_{name}"]
    if percent:
        return f"{100 * mean_value:.2f}±{100 * sd_value:.2f}%"
    return f"{mean_value:.6f}±{sd_value:.6f}"


def _add_partition_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = True
    for cell, text in zip(table.rows[0].cells, ("子集", "图像数", "用途", "是否参与专家梯度")):
        _set_cell(cell, text, True, 9.0, "D9E2F3")
    _repeat_header(table.rows[0])
    rows = (
        ("D_E：专家拟合", "4,176", "训练共享骨干、双专家与校验器", "是"),
        ("D_S：专家停止", "893", "早停与专家检查点选择", "否"),
        ("D_G：门控未见", "892", "样本外最终风险监督", "否"),
        ("D_seen：配对对照", "892", "专家已见样本上的门控监督", "已参与"),
    )
    for values in rows:
        row = table.add_row()
        for cell, text in zip(row.cells, values):
            _set_cell(cell, text, False, 8.8)


def _add_results_table(document: Document, analysis: dict) -> None:
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.autofit = True
    headers = ("策略", "最终 NLL/nat", "Macro F1", "目标召回", "含钛误入", "金属误入", "平均门控")
    for cell, text in zip(table.rows[0].cells, headers):
        _set_cell(cell, text, True, 7.8, "D9E2F3")
    _repeat_header(table.rows[0])

    equal = analysis["references"]["equal"]
    original = analysis["references"]["original_joint_gate"]
    rows = [
        (
            "等权融合",
            f"{equal['final_nll_nats']:.6f}",
            f"{equal['final_macro_f1']:.4f}",
            f"{100 * equal['final_target_recall']:.2f}%",
            f"{100 * equal['final_ti_intrusion']:.2f}%",
            f"{100 * equal['final_metal_intrusion']:.2f}%",
            f"{equal['mean_gate']:.3f}",
        ),
        (
            "原联合门控",
            f"{original['final_nll_nats']:.6f}",
            f"{original['final_macro_f1']:.4f}",
            f"{100 * original['final_target_recall']:.2f}%",
            f"{100 * original['final_ti_intrusion']:.2f}%",
            f"{100 * original['final_metal_intrusion']:.2f}%",
            f"{original['mean_gate']:.3f}",
        ),
    ]
    aggregates = _aggregate_map(analysis)
    labels = (
        ("seen + 最终NLL", "seen", "final_nll"),
        ("unseen + 最终NLL", "unseen", "final_nll"),
        ("seen + NLL/遗憾", "seen", "final_nll_plus_regret"),
        ("unseen + NLL/遗憾", "unseen", "final_nll_plus_regret"),
    )
    for label, source, objective in labels:
        row = aggregates[(source, objective)]
        rows.append((
            label,
            _metric(row, "final_nll_nats"),
            _metric(row, "final_macro_f1"),
            _metric(row, "final_target_recall", True),
            _metric(row, "final_ti_intrusion", True),
            _metric(row, "final_metal_intrusion", True),
            _metric(row, "mean_gate"),
        ))
    for values in rows:
        row = table.add_row()
        for cell, text in zip(row.cells, values):
            _set_cell(cell, text, False, 7.1)


def append_appendix(report: Path, figure: Path, analysis: dict, degeneracy: dict) -> bool:
    if not figure.exists():
        raise FileNotFoundError(figure)
    if analysis.get("independence_warning") is None:
        raise ValueError("Missing independence warning")
    if abs(degeneracy["equal_experts"]["final_nll_gradient"]) > 1e-12:
        raise ValueError("Degeneracy verification did not pass")

    document = Document(report)
    if any(paragraph.text == TITLE for paragraph in document.paragraphs):
        return False

    before_images = len(document.inline_shapes)
    before_tables = len(document.tables)
    _keep_phr_result_header_with_data(document)
    document.add_page_break()
    _add_heading(document, TITLE, 1)
    _add_body(
        document,
        "本附录在 RSG-HRGV 的共享骨干、角色/矿种双专家、固定种类到角色映射以及两类困难负样本校验器基础上，提出样本外风险监督路由 OOS-RSG-HRGV。其网络改进不依赖继续堆叠普通注意力模块，而是改变门控的学习位置和目标：专家在一组数据上拟合并停止，门控在按原始图像组隔离的未见样本上直接最小化最终校验后风险。当前结果用于验证这一机制及其理论边界，尚不构成独立外部泛化或稳定优越性结论。",
    )

    _add_heading(document, "K.1 数据隔离与网络结构", 2)
    _add_body(
        document,
        "为区分专家拟合能力与门控泛化能力，数据先按 photo ID、重复/近重复组进行不可拆分分组，再构造四个功能子集。D_G 与 D_seen 在 17 个矿种-角色分层上的计数逐项相同，且 D_G 不与 D_E、D_S 或 D_seen 共享图像组。钛磁铁矿在完整数据中仍为 35 张；本实验两个配对子集各取 2 张只是受严格分层匹配约束，并不表示全数据只有 2 张或 23 张。",
    )
    _add_partition_table(document)
    _add_equation(document, "θ̂ = A(D_E, D_S),    ∂θ̂/∂L_gate = 0", "K-1")
    _add_body(
        document,
        "式（K-1）表示专家参数由拟合集和停止集确定，进入门控训练后全部参数与缓冲区冻结。共享 EfficientNet-B0 输出 1,280 维特征 h；直接角色头给出 p_d，矿种头给出 p_s，经固定映射矩阵 A 得到 p_m=A p_s；含钛与金属光泽校验器输出 v_Ti、v_M。门控输入同时包含视觉特征、两专家后验、熵、JS 分歧和校验器证据。",
    )
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(figure), width=Cm(15.4))
    caption = document.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    _set_run_font(caption.add_run("图 28 OOS-RSG-HRGV 的数据隔离、双专家结构与样本外风险监督路径"), "宋体", 10.5)

    _add_heading(document, "K.2 最终风险监督与网络训练目标", 2)
    _add_body(
        document,
        "设 g_φ(u) 为样本相关门控，p_d 与 p_m 分别为直接角色专家和矿种映射专家的角色后验，V 为已冻结的校验修正算子。与只拟合标签相关软 oracle 的原门控不同，OOS-RSG 首先在最终决策分布上定义风险。",
    )
    _add_equation(document, "p_g(x)=g_φ(u)p_d(x)+[1−g_φ(u)]p_m(x)", "K-2")
    _add_equation(document, "p_f(x)=V[p_g(x),v_Ti(x),v_M(x)]", "K-3")
    _add_equation(document, "L_OOS(φ)=E_(x,y)∼D_G[−log p_f(y|x)] + λ E_(x,y)∼D_G[BCE(g_φ,t_regret)]", "K-4")
    _add_body(
        document,
        "其中 λ=0 对应纯最终 NLL；λ=0.1 对应最终 NLL 加遗憾辅助监督。t_regret 只作为低权重归纳偏置，最终风险仍通过校验后的 p_f 回传至门控。该设计把网络创新落实为“校验器感知、样本外、最终风险监督”的路由，而不是把数据划分技术本身宣称为新算法。",
    )

    _add_heading(document, "K.3 命题一：固定门控的条件无偏风险与梯度", 2)
    _add_body(
        document,
        "命题 1：条件于已由 D_E、D_S 决定的冻结专家 θ̂，若 D_G 由与专家拟合组相互独立的图像组从目标分布 P 抽取，且待评价的门控参数 φ 在观察 D_G 前固定，则 D_G 上的经验最终风险及其梯度分别是条件总体风险与梯度的无偏估计：",
    )
    _add_equation(document, "E[ R̂_G(φ) | D_E,D_S ] = R(φ | θ̂)", "K-5")
    _add_equation(document, "E[ ∇_φ R̂_G(φ) | D_E,D_S ] = ∇_φ R(φ | θ̂)", "K-6")
    _add_body(
        document,
        "证明要点：在有限期望及可交换微分与期望的常规条件下，独立同分布样本损失和梯度的样本均值分别对总体期望无偏。其意义是：与在 D_E 上给门控制造标签相比，D_G 能提供未参与专家拟合的数据上的风险信号。但当 φ 已在同一 D_G 上优化后，再用 D_G 报告其泛化风险并不保持该无偏结论；本实验还复用了 D_S 选门控轮次，因此 D_S 不是最终独立确认集。",
    )

    _add_heading(document, "K.4 命题二：双专家一致时的门控退化", 2)
    _add_body(
        document,
        "命题 2：若某输入上 p_d(x)=p_m(x)，则任意 g_φ 都有 p_g(x)=p_d(x)=p_m(x)。当校验器只依赖 p_g 与冻结校验证据时，p_f 也与门控无关，最终 NLL 对门控参数的梯度为零；同时两专家真实类损失差为零，基于专家差距的 RSG 权重也为零。",
    )
    _add_equation(document, "p_d=p_m  ⇒  ∂[−log p_f(y|x)]/∂φ = 0", "K-7")
    _add_equation(document, "|ℓ_d−ℓ_m|=0  ⇒  w_gap=tanh(|ℓ_d−ℓ_m|/τ_w)=0", "K-8")
    _add_body(
        document,
        f"float64 数值核验中，相同专家的最终 NLL 门控梯度为 {degeneracy['equal_experts']['final_nll_gradient']:.1f}，差距权重为 {degeneracy['equal_experts']['gap_weight']:.1f}；构造不同专家后，梯度为 {degeneracy['distinct_experts']['final_nll_gradient']:.6f}，差距权重为 {degeneracy['distinct_experts']['gap_weight']:.6f}。该核验确认公式和实现代数一致，但不是性能证据。该命题揭示了路由网络的必要条件：只有双专家在样本上提供不同预测信息，门控才具有可学习信号。",
    )

    _add_heading(document, "K.5 seen/unseen 配对门控实验", 2)
    _add_body(
        document,
        "固定一个完整训练的 HRGV 专家，在完全相同的三组门控初始化下，分别使用专家已见的 D_seen 与组隔离的 D_G 训练门控；每组运行 30 轮，并在同一 893 张停止集上按最低最终 NLL 选择轮次。全部 12 次运行均保持专家参数、缓冲区逐项不变且无专家梯度，未访问锁定测试集。表中学习门控为三组配对初始化的均值±样本标准差；等权融合与原联合门控是同一冻结专家上的确定性参考。",
    )
    _add_results_table(document, analysis)
    paired = {row["objective"]: row for row in analysis["paired_aggregates"]}
    nll_pair = paired["final_nll"]
    hybrid_pair = paired["final_nll_plus_regret"]
    _add_body(
        document,
        f"在纯最终 NLL 目标下，unseen 相对 seen 的配对差为：最终 NLL {nll_pair['mean_unseen_minus_seen_final_nll_nats']:.6f} nat、Macro F1 +{nll_pair['mean_unseen_minus_seen_final_macro_f1']:.6f}、目标召回 +{100*nll_pair['mean_unseen_minus_seen_final_target_recall']:.2f} 个百分点；三组初始化方向一致。金属误入率平均增加 {100*nll_pair['mean_unseen_minus_seen_final_metal_intrusion']:.2f} 个百分点，说明改进并非所有指标共同占优。加入遗憾辅助项后，unseen 相对 seen 的 NLL 差仅 {hybrid_pair['mean_unseen_minus_seen_final_nll_nats']:.6f} nat，目标召回差为零，监督源优势明显减弱。",
    )
    _add_body(
        document,
        "相较等权融合，unseen + NLL/遗憾在本停止集上的最终 NLL 与 Macro F1 略优，但目标召回更低；纯 unseen NLL 也未同时占优于等权融合。因此现阶段能够支持的结论是：专家已见样本会使门控监督表现出明显拟合乐观性，组隔离的样本外最终风险能够改变并改善部分配对指标；尚不能宣称 OOS-RSG-HRGV 已稳定优于简单等权融合或原主模型。",
    )

    _add_heading(document, "K.6 理论创新定位与后续确认", 2)
    _add_body(
        document,
        "本附录形成的候选网络创新由四个相互约束的部分组成：（1）角色专家与矿种映射专家构成具有领域语义的异构双专家；（2）含钛与金属光泽校验器将困难负样本知识写入最终后验；（3）门控不再在专家拟合样本上学习，而在图像组隔离的未见样本上优化校验后最终风险；（4）以条件无偏性和门控退化命题限定方法何时具有有效监督。样本拆分、交叉拟合和代价敏感思想本身均为已有方法，本研究不将其单独表述为首创。",
    )
    _add_body(
        document,
        "当前三组门控随机种子共享同一个冻结专家和同一停止集，所以不是三个独立专家，也不是独立确认。论文中若要把 OOS-RSG-HRGV 从候选网络提升为正式主方法，下一步必须训练至少三组独立专家种子，为每组专家重新构造门控监督，并在未参与专家早停、门控训练或轮次选择的最终评价集上检验；同时报告最终 NLL、Macro F1、目标召回、含钛误入和金属误入。来源外矿区/摄影者数据和真实矿石图像应作为更高层级泛化证据，不能由当前公开标本结果替代。",
    )
    _add_body(
        document,
        "复现文件：scripts/run_seen_unseen_gate_study.py、scripts/analyze_seen_unseen_gate_study.py、scripts/verify_oos_gate_degeneracy.py；结果位于 outputs/training/seen_unseen_gate_study_v1 与 outputs/theory/oos_gate_degeneracy.json。",
    )

    temporary = report.with_suffix(".oos_rsg.tmp.docx")
    document.save(temporary)
    checked = Document(temporary)
    if len(checked.inline_shapes) != before_images + 1:
        raise AssertionError("Expected exactly one appended figure")
    if len(checked.tables) != before_tables + 2:
        raise AssertionError("Expected exactly two appended tables")
    if sum(paragraph.text == TITLE for paragraph in checked.paragraphs) != 1:
        raise AssertionError("Appendix title missing or duplicated")
    temporary.replace(report)
    return True


def main() -> None:
    if not REPORT.exists():
        raise FileNotFoundError(REPORT)
    for path in (FIGURE, ANALYSIS, DEGENERACY):
        if not path.exists():
            raise FileNotFoundError(path)
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists():
        shutil.copy2(REPORT, BACKUP)
    analysis = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    degeneracy = json.loads(DEGENERACY.read_text(encoding="utf-8"))
    changed = append_appendix(REPORT, FIGURE, analysis, degeneracy)
    print("updated" if changed else "already_present")


if __name__ == "__main__":
    main()
