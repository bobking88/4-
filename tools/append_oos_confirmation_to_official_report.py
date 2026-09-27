"""Append the independent OOS-RSG confirmation and constrained-router theory."""
from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from append_oos_rsg_to_official_report import (
    _add_body,
    _add_equation,
    _add_heading,
    _repeat_header,
    _set_cell,
    _set_run_font,
)


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "结题" / "基于深度学习的钒钛矿相关矿物图像识别方法研究_技术报告（正式版）.docx"
FIGURE = ROOT / "outputs" / "paper_figures_v3" / "fig_oos_rsg_confirmation.png"
ANALYSIS = ROOT / "outputs" / "training" / "oos_rsg_confirmation_v1" / "analysis" / "analysis.json"
AUDIT = ROOT / "outputs" / "training" / "oos_rsg_confirmation_manifests_v1" / "audit.json"
BACKUP = ROOT / "结题" / "历史版本" / "基于深度学习的钒钛矿相关矿物图像识别方法研究_技术报告（正式版_20260924_追加独立确认前）.docx"
TITLE = "附录 L 多专家样本外路由确认与目标召回约束"


def _repair_existing_appendix(document: Document) -> bool:
    old = "表 L-2 中效应统一定义为"
    new = "下表中效应统一定义为"
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                return True
    return False


def _add_protocol_table(document: Document, audit: dict) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ("功能子集", "图像数", "参数或选择用途", "与 final_eval 关系")
    for cell, text in zip(table.rows[0].cells, headers):
        _set_cell(cell, text, True, 8.5, "D9E2F3")
    _repeat_header(table.rows[0])
    counts = audit["partition_counts"]
    rows = (
        ("expert_fit", counts["expert_fit"], "专家参数更新", "组隔离"),
        ("expert_stop", counts["expert_stop"], "专家轮次选择", "组隔离"),
        ("gate_seen", audit["counts"]["gate_seen"], "专家训练内监督对照", "组隔离"),
        ("gate_unseen", audit["counts"]["gate_unseen"], "样本外门控拟合", "组隔离"),
        ("gate_stop", counts["gate_stop"], "门控轮次选择", "组隔离"),
        ("final_eval", counts["final_eval"], "全部选择锁定后只读评价", "自身"),
    )
    for values in rows:
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            _set_cell(cell, str(value), False, 8.2)


def _effect_text(summary: dict, *, percent: bool) -> str:
    scale = 100.0 if percent else 1.0
    suffix = " 点" if percent else ""
    return (
        f"{scale * summary['difference']:+.3f}"
        f" [{scale * summary['ci_low']:+.3f}, {scale * summary['ci_high']:+.3f}]{suffix}"
    )


def _add_effect_table(document: Document, analysis: dict) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ("指标", "unseen - seen", "unseen - 等权", "方向解释")
    for cell, text in zip(table.rows[0].cells, headers):
        _set_cell(cell, text, True, 8.2, "D9E2F3")
    _repeat_header(table.rows[0])
    seen = analysis["comparisons"]["unseen_minus_seen"]["summary"]
    equal = analysis["comparisons"]["unseen_minus_equal"]["summary"]
    specs = (
        ("最终 NLL / nat", "final_nll_nats", False, "负值有利"),
        ("Accuracy", "accuracy", True, "正值有利"),
        ("Macro F1", "macro_f1", True, "正值有利"),
        ("目标召回", "target_recall", True, "正值有利"),
        ("含钛误入目标", "ti_intrusion", True, "负值有利"),
        ("金属光泽误入目标", "metal_intrusion", True, "负值有利"),
    )
    for label, key, percent, direction in specs:
        row = table.add_row()
        values = (
            label,
            _effect_text(seen[key], percent=percent),
            _effect_text(equal[key], percent=percent),
            direction,
        )
        for cell, text in zip(row.cells, values):
            _set_cell(cell, text, False, 7.7)


def _add_decision_table(document: Document, analysis: dict) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ("预注册判据", "结果", "判定")
    for cell, text in zip(table.rows[0].cells, headers):
        _set_cell(cell, text, True, 8.5, "D9E2F3")
    _repeat_header(table.rows[0])
    criteria = analysis["promotion_decision"]["criteria"]
    rows = (
        ("等权对照 NLL 区间上界小于 0", "满足", criteria["nll_ci_below_zero"]),
        ("至少 2/3 专家 NLL 方向有利", "3/3", criteria["nll_favorable_in_at_least_two_seeds"]),
        ("Macro F1 下降不超过 0.5 点", "下降 0.047 点", criteria["macro_f1_guardrail"]),
        ("目标召回下降不超过 1 点", "下降 2.476 点", criteria["target_recall_guardrail"]),
        ("含钛误入上升不超过 1 点", "下降 1.170 点", criteria["ti_intrusion_guardrail"]),
        ("金属误入上升不超过 1 点", "下降 1.600 点", criteria["metal_intrusion_guardrail"]),
    )
    for label, result, passed in rows:
        row = table.add_row()
        values = (label, result, "通过" if passed else "未通过")
        for index, (cell, text) in enumerate(zip(row.cells, values)):
            shade = "FCE4D6" if index == 2 and not passed else None
            _set_cell(cell, text, index == 2, 8.2, shade)


def append_appendix(report: Path, figure: Path, analysis: dict, audit: dict) -> bool:
    if not figure.exists():
        raise FileNotFoundError(figure)
    if not audit.get("final_eval_locked"):
        raise ValueError("final_eval was not locked")
    if not audit.get("original_validation_and_test_omitted"):
        raise ValueError("original validation/test were not omitted")
    if analysis["promotion_decision"]["promote_to_main_method"]:
        raise ValueError("Expected the preregistered promotion decision to fail")

    document = Document(report)
    if any(paragraph.text == TITLE for paragraph in document.paragraphs):
        if not _repair_existing_appendix(document):
            return False
        temporary = report.with_suffix(".oos_confirmation.tmp.docx")
        document.save(temporary)
        temporary.replace(report)
        return True

    before_images = len(document.inline_shapes)
    before_tables = len(document.tables)
    document.add_page_break()
    _add_heading(document, TITLE, 1)
    _add_body(
        document,
        "附录 K 的单专家探索表明，组隔离的专家未见图像可能改善门控的最终概率风险，但三组门控初始化共享同一专家与停止集，不能承担独立确认。本附录按预先冻结的五段式协议重新训练三位独立专家，在全部专家与门控选择完成后只读最终评价集，并据此决定候选网络是否晋级。",
    )

    _add_heading(document, "L.1 五段式隔离协议与独立专家", 2)
    _add_body(
        document,
        "实验只重划原固定训练集，重复和近重复图像由 split_group_id 保持不可拆分，原验证集和原测试集均不参与。gate_seen 与 gate_unseen 在 17 个矿种-角色分层上的计数逐项相同；除 gate_seen 按设计属于 expert_fit 外，其余功能子集之间不存在图像组交叉。完整数据集仍有 35 张钛磁铁矿图像，本协议的 final_eval 仅含 2 张，不能据此报告稳定的钛磁铁矿单类性能。",
    )
    _add_protocol_table(document, audit)
    _add_body(
        document,
        "三位 EfficientNet-B0 专家分别以 20260924、20260925 和 20260926 为种子，从 ImageNet 权重独立初始化。每位专家均冻结后训练一对 seen/unseen 门控，两者共享门控初始化、批次顺序和超参数，仅监督来源不同。门控分别按独立 gate_stop 的最终 NLL 选轮次，专家参数和缓冲区在门控训练中逐项保持不变。",
    )
    _add_equation(document, "θ̂_j=A(D_E,D_S; seed_j),    j=1,2,3", "L-1")
    _add_equation(document, "φ̂_j=arg min_φ R̂_(G,j)(φ | θ̂_j)", "L-2")

    _add_heading(document, "L.2 独立确认结果与错误类型转移", 2)
    _add_body(
        document,
        "主要终点为 final_eval 上的校验后最终 NLL。下表中效应统一定义为 unseen 减对照；区间来自 10,000 次两阶段成对 Bootstrap，先重采样专家种子，再在各矿种-角色层内按 split_group_id 重采样图像组。NLL 和误入率为负值有利，Accuracy、Macro F1 和目标召回为正值有利。",
    )
    _add_effect_table(document, analysis)
    _add_body(
        document,
        "相对等权融合，样本外门控在三位专家上均降低最终 NLL，平均差为 -0.014040 nat，95% 区间为 [-0.019431, -0.008724] nat；含钛与金属光泽误入目标平均分别下降 1.170 和 1.600 个百分点。目标召回则平均下降 2.476 个百分点，95% 区间为 [-4.381, -0.571] 个百分点。结果说明样本外监督稳定改善了正确概率风险并减少困难负样本误入，但其风险重分配以目标矿物漏识增加为代价。",
    )
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(figure), width=Cm(16.2))
    caption = document.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    _set_run_font(
        caption.add_run(
            "图 29 三专家样本外风险路由的独立确认效应与预注册晋级判定"
        ),
        "宋体",
        10.5,
    )

    _add_heading(document, "L.3 预注册晋级判定", 2)
    _add_decision_table(document, analysis)
    _add_body(
        document,
        "预注册规则要求主要 NLL 终点、跨专家方向和四项伤害护栏全部通过。目标召回下降超过 1 个百分点上限，因此 OOS-RSG-HRGV 不晋级为第一篇论文主方法。该负结果不否定样本外风险监督的机制价值，而是把下一步网络问题从无约束门控收缩为具有目标保护条件的风险路由。",
    )

    _add_heading(document, "L.4 目标召回约束的样本外风险路由", 2)
    _add_body(
        document,
        "下一候选网络暂记为 TC-OOS-RSG。固定双专家与校验器后，它在样本外数据上最小化最终 NLL，同时要求目标漏识风险不超过等权融合基线加容许量 ε。该问题可写为：",
    )
    _add_equation(
        document,
        "min_φ R_NLL(φ)    s.t.    R_T,miss(φ) ≤ R_T,miss(g≡1/2)+ε",
        "L-3",
    )
    _add_equation(
        document,
        "L(φ,λ)=R_NLL(φ)+λ[R_T,miss(φ)−τ],    λ≥0",
        "L-4",
    )
    _add_body(
        document,
        "为保留可审计回退，可在校验后把等权后验 q_1/2 与无约束样本外后验 q_φ 做最终概率插值。令 0≤ρ≤1：",
    )
    _add_equation(document, "q_ρ=(1−ρ)q_1/2+ρq_φ", "L-5")
    _add_body(
        document,
        "定义目标软漏识风险 S_T(q)=E[1−q_T(X) | Y=T]。该风险沿插值路径为仿射函数，而 NLL 由负对数的凸性得到凸组合上界：",
    )
    _add_equation(
        document,
        "S_T(q_ρ)=S_T(q_1/2)+ρ[S_T(q_φ)−S_T(q_1/2)]",
        "L-6",
    )
    _add_equation(
        document,
        "R_NLL(q_ρ)≤(1−ρ)R_NLL(q_1/2)+ρR_NLL(q_φ)",
        "L-7",
    )
    _add_equation(
        document,
        "ρ≤min{1, ε/[S_T(q_φ)−S_T(q_1/2)]},    if S_T(q_φ)>S_T(q_1/2)",
        "L-8",
    )
    _add_body(
        document,
        "上述关系给出风险改善与目标保护之间可计算的连续路径，且 ρ=0 始终保留等权回退。但软概率风险约束不直接等同于离散 argmax 召回保证，仍需独立约束集、有限样本安全裕量和新的外层评价协议。现有 final_eval 已经被读取，不能再用于 TC-OOS-RSG 的确认性调参或最终确认；本节因此属于由独立失效模式驱动的下一理论假设，而不是已完成的性能贡献。",
    )

    _add_heading(document, "L.5 结论边界与复现文件", 2)
    _add_body(
        document,
        "独立确认支持的结论是：对公开矿物标本图像中的跨粒度双专家，样本外最终风险监督能够稳定降低最终 NLL，并减少含钛和金属光泽困难负样本误入；它未保留目标召回，故不能写作全面优于等权融合。该结论不等同于真实矿石颗粒、工业皮带、品位、回收率、元素含量或跨网站验证。",
    )
    _add_body(
        document,
        "复现文件包括 scripts/build_oos_confirmation_manifests.py、scripts/run_oos_confirmation_experts.py、scripts/run_oos_rsg_confirmation.py、scripts/analyze_oos_rsg_confirmation.py 和 scripts/generate_oos_confirmation_figure.py；协议、逐图预测、统计区间和晋级判定保存在 outputs/training/oos_rsg_confirmation_manifests_v1 与 outputs/training/oos_rsg_confirmation_v1。",
    )

    temporary = report.with_suffix(".oos_confirmation.tmp.docx")
    document.save(temporary)
    checked = Document(temporary)
    if len(checked.inline_shapes) != before_images + 1:
        raise AssertionError("Expected exactly one appended confirmation figure")
    if len(checked.tables) != before_tables + 3:
        raise AssertionError("Expected exactly three appended confirmation tables")
    if sum(paragraph.text == TITLE for paragraph in checked.paragraphs) != 1:
        raise AssertionError("Appendix title missing or duplicated")
    temporary.replace(report)
    return True


def main() -> None:
    for path in (REPORT, FIGURE, ANALYSIS, AUDIT):
        if not path.exists():
            raise FileNotFoundError(path)
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists():
        shutil.copy2(REPORT, BACKUP)
    analysis = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    audit = json.loads(AUDIT.read_text(encoding="utf-8"))
    changed = append_appendix(REPORT, FIGURE, analysis, audit)
    print("updated" if changed else "already_present")


if __name__ == "__main__":
    main()
