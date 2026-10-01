"""Append auditable ABMP derivations and development evidence to the sole report."""

from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

from append_oos_rsg_to_official_report import _add_body, _add_heading as _existing_heading, _repeat_header, _set_cell, _set_run_font


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "结题" / "基于深度学习的钒钛矿相关矿物图像识别方法研究_技术报告（正式版）.docx"
TITLE = "附录 M 自适应预算与目标间隔保护：理论推导及开发验证"
OWNER_PREFIX = "ABMP_M_"


def _appendix_nodes(document, title_node):
    children = list(document._element.body)
    return [node for node in children[children.index(title_node):] if node.tag != qn("w:sectPr")]


def _appendix_digest(nodes, part):
    copies = [deepcopy(node) for node in nodes]
    owner_ids = set()
    for node in copies:
        for marker in list(node.iter(qn("w:bookmarkStart"))):
            if marker.get(qn("w:name"), "").startswith(OWNER_PREFIX):
                owner_ids.add(marker.get(qn("w:id")))
                marker.getparent().remove(marker)
    for node in copies:
        for marker in list(node.iter(qn("w:bookmarkEnd"))):
            if marker.get(qn("w:id")) in owner_ids:
                marker.getparent().remove(marker)
    references = set()
    relationship_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for node in copies:
        for element in node.iter():
            references.update(value for key, value in element.attrib.items() if key.startswith(relationship_namespace))
    relationships = []
    for reference in sorted(references):
        if reference not in part.rels:
            raise ValueError("Appendix has a missing package relationship.")
        relation = part.rels[reference]
        blob_hash = None if relation.is_external else hashlib.sha256(relation.target_part.blob).hexdigest()
        relationships.append([reference, relation.reltype, relation.target_ref, blob_hash])
    content = b"".join(etree.tostring(node, method="c14n") for node in copies)
    content += json.dumps(relationships, ensure_ascii=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(content).hexdigest()[:24]


def _mark_owned_appendix(document, title_node):
    markers = [node for node in title_node.iter(qn("w:bookmarkStart")) if node.get(qn("w:name"), "").startswith(OWNER_PREFIX)]
    old_ids = {node.get(qn("w:id")) for node in markers}
    for marker in markers + [node for node in title_node.iter(qn("w:bookmarkEnd")) if node.get(qn("w:id")) in old_ids]:
        marker.getparent().remove(marker)
    digest = _appendix_digest(_appendix_nodes(document, title_node), document.part)
    existing_ids = [int(node.get(qn("w:id"))) for node in document._element.iter(qn("w:bookmarkStart"))]
    marker_id = str(max(existing_ids, default=0) + 1)
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), marker_id)
    start.set(qn("w:name"), OWNER_PREFIX + digest)
    end.set(qn("w:id"), marker_id)
    title_node.insert(1 if title_node.find(qn("w:pPr")) is not None else 0, start)
    title_node.append(end)


def _is_standalone_page_break(node):
    if node is None or node.tag != qn("w:p") or len(node) != 1:
        return False
    run = node[0]
    return (run.tag == qn("w:r") and len(run) == 1
            and run[0].tag == qn("w:br") and run[0].get(qn("w:type")) == "page")


def _add_heading(document, text, level):
    _existing_heading(document, text, level)
    for run in document.paragraphs[-1].runs:
        run.font.color.rgb = RGBColor(0, 0, 0)


def render_equations(folder: Path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    folder.mkdir(parents=True, exist_ok=True)
    formulas = {
        "policy": r"$h_i=F_\psi(z_i),\quad\widetilde\rho_i=\sigma(w_r^Th_i+b_r),\quad\varepsilon_i=\varepsilon_{max}\sigma(w_\varepsilon^Th_i+b_\varepsilon-\lambda)$",
        "loss": r"$\mathcal{L}=-\log q_{i,y_i}+0.25\,\mathrm{BCE}(h_r,u_i)+0.10\,\mathrm{BCE}_{d_i>0}(h_\varepsilon,v_i)$",
        "targets": r"$u_i=\sigma\left(\frac{\log q_{\phi,y_i}-\log q_{0,y_i}}{0.25}\right),\quad v_i=u_i\min\left(1,\frac{(q_{0,T}-q_{\phi,T})_+}{\varepsilon_{max}}\right)$",
        "posterior": r"$d_i=(q_{0,T}-q_{\phi,T})_+,\quad c_i^{post}=\min(1,\varepsilon_i/d_i)\ (d_i>0),\quad q_T\geq q_{0,T}-\varepsilon_i$",
        "margin": r"$\Delta_{i,k}=(q_{\phi,k}-q_{0,k})-(q_{\phi,T}-q_{0,T}),\quad c_i^{margin}=\min_{k:\Delta_{i,k}>0}\mathrm{clip}_{[0,1]}\left(\frac{q_{0,T}-q_{0,k}-\delta}{\Delta_{i,k}}\right)$",
        "fusion": r"$\widehat\rho_i=\min(\widetilde\rho_i,c_i^{post},c_i^{margin}),\quad q_i=(1-\widehat\rho_i)q_{0,i}+\widehat\rho_iq_{\phi,i}$",
        "projection": r"$\widehat\rho_i=\arg\min_{r\in[0,\min(c_i^{post},c_i^{margin})]}\frac{1}{2}(r-\widetilde\rho_i)^2,\quad A_i=1\Rightarrow q_T-q_k\geq\delta>0$",
        "dominance": r"$q_T-\max_{k\ne T}q_k\geq 2q_T-1\geq 2(q_{0,T}-\varepsilon_i)-1,\quad\tau_p-\varepsilon_{max}>(1+\delta)/2$",
        "capacity": r"$D_i=\mathbf{1}[\arg\max q_{0,i}\ne\arg\max q_{\phi,i}],\quad |\mathrm{Acc}(q)-\mathrm{Acc}(q_0)|\leq\frac{1}{N}\sum_iD_i$",
        "recall": r"$|\mathrm{Rec}_T(q)-\mathrm{Rec}_T(q_0)|\leq\frac{\sum_{i:y_i=T}D_i}{N_T},\quad\mathrm{Rec}_T(q)-\mathrm{Rec}_T(q_0)\leq\frac{\sum_{i:y_i=T,\hat y_{0,i}\ne T}D_i}{N_T}$",
    }
    paths = {}
    for name, formula in formulas.items():
        fig = plt.figure(figsize=(14, .75), facecolor="white")
        artist = fig.text(.5, .5, formula, ha="center", va="center", fontsize=16, color="black")
        fig.canvas.draw()
        width = artist.get_window_extent(fig.canvas.get_renderer()).width / fig.dpi
        if width > 13.5:
            artist.set_fontsize(16 * 13.5 / width)
        path = folder / f"abmp_{name}.png"
        fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=.13)
        plt.close(fig)
        paths[name] = path
    return paths


def _equation(document, path, number):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Cm(15.8))
    label = document.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label.paragraph_format.space_after = Pt(4)
    _set_run_font(label.add_run(f"({number})"), "Times New Roman", 10)


def _table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    _repeat_header(table.rows[0])
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for index, values in enumerate([headers, *rows]):
        row = table.rows[0] if index == 0 else table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            _set_cell(cell, str(value), index == 0, 9 if len(headers) < 5 else 8.5, "E7EBEE" if index == 0 else None)
            if index == 0:
                cell.paragraphs[0].paragraph_format.keep_with_next = True
    document.add_paragraph()


def _percent(value):
    return f"{100 * value:.2f}%"


def append_appendix(report: Path, summaries: list, geometry: dict, invariants: dict, figure_dir: Path, *, refresh=False):
    if len(summaries) != 2 or any(s.get("purpose") != "development_only" for s in summaries):
        raise ValueError("Only the two exposed Fold 0 development rounds may be reported here.")
    r1, r2 = summaries
    if r2["status"] != "DEVELOPMENT_GATE_FAILED" or r2["selected"] is not None:
        raise ValueError("This appendix requires the observed unsuccessful r2 evidence.")
    if geometry["source_status"] != r2["status"] or invariants["purpose"] != "formula_verification_not_classification_performance":
        raise ValueError("Evidence purpose mismatch.")
    figure = figure_dir / "fig_abmp_rsg_architecture.png"
    if not figure.is_file():
        raise FileNotFoundError(figure)
    document = Document(report)
    previous = next((p for p in document.paragraphs if p.text == TITLE), None)
    if previous is not None:
        if not refresh:
            return False
        markers = [marker for marker in previous._p.iter(qn("w:bookmarkStart")) if marker.get(qn("w:name"), "").startswith(OWNER_PREFIX)]
        nodes = _appendix_nodes(document, previous._p)
        if len(markers) != 1 or markers[0].get(qn("w:name")) != OWNER_PREFIX + _appendix_digest(nodes, document.part):
            raise ValueError("Appendix ownership or content changed; refusing to remove user content.")
        preceding = previous._p.getprevious()
        body = document._element.body
        for child in nodes:
            body.remove(child)
        if _is_standalone_page_break(preceding):
            body.remove(preceding)
    formulas = render_equations(figure_dir / "formulas")
    document.add_page_break()
    _add_heading(document, TITLE, 1)
    title_node = document.paragraphs[-1]._p
    _add_body(document, "本附录将网络的解析性质、数学推导和真实开发实验分开呈现。ABMP-RSG-Net v2 已实现双头预算策略及双重投影，但第二轮开发门未通过，不能替代已确认的基线，也不能表述为性能创新已成立。Fold 0 已用于方法开发，以下结果不是独立确证；Fold 1/2 未用于本轮评价，保留至新方案冻结后使用。")

    _add_heading(document, "M.1 双头网络结构与证据", 2)
    _add_body(document, "冻结 EfficientNet-B0、直接角色头、矿物种类头、种类-角色映射、含钛及金属光泽校验器，并复用样本外训练的候选门。校验器完整的 q0 和 qφ 构成两端后验；18 维证据包括两路后验、分歧、熵及校验信息。仅训练 2,418 参数的预算-路由策略模块，推理时单图预算不依赖其他批次样本。")
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(str(figure), width=Cm(16.0))
    caption = document.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(caption.add_run("图 M-1 ABMP-RSG-Net v2：冻结专家、可训练双头策略与解析投影"), "宋体", 10)
    caption.runs[-1].font.color.rgb = RGBColor(0, 0, 0)
    _table(document, ["组件", "结构/作用", "本轮是否更新"], [
        ["图像专家及校验器", "EfficientNet-B0 + 双头 + 两类校验器", "否"],
        ["候选门", "组隔离、样本外监督训练的门控", "否"],
        ["共享策略主干", "18 → 64 → LN → SiLU → Dropout(0.10) → 16 → SiLU", "是"],
        ["路由/预算双头", "两个独立的 16 → 1 线性层 + Sigmoid", "是"],
        ["后验/间隔投影", "解析上限取最小；无可训练参数", "否"],
    ], [3.2, 10.2, 2.6])
    _add_heading(document, "M.2 问题定义、可证伪假设与训练目标", 2)
    _add_body(document, "固定风险预算不激活时，约束结构可能只是无效附加层；高置信度目标锚点也可能为假阳性。研究问题因此不是增加网络层数，而是在候选专家具有互补性的前提下，以可学习的逐图预算调节采用幅度，再用解析投影限制目标后验下降并保持指定锚点的类别间隔。待检验假设是：预算分配必须优于同预算对照，间隔层必须有独立作用，风险改善不能以超出协议的目标召回下降为代价。")
    _equation(document, formulas["policy"], "M-1")
    _add_body(document, "其中 z_i 是 18 维证据，λ 为开发集固定偏移，εmax 为允许的最大后验预算。双头共享表示但不共享输出层；预算范围来自 Sigmoid 与上限缩放，而不是按测试批次重新归一化。")
    _equation(document, formulas["loss"], "M-2")
    _equation(document, formulas["targets"], "M-3")
    _add_body(document, "u_i 由候选后验与回退后验的真实类对数概率差构造；v_i 同时考虑收益和目标后验下降。预算监督仅在 d_i>0 的拟合图像上启用。停止轮次按独立 projector_stop 的最终 NLL 决定。BCE 采用 logits 实现；这些软目标是当前可检验设计，不是天然最优预算的真值。")

    _add_heading(document, "M.3 解析投影、可行域与目标保持", 2)
    _equation(document, formulas["posterior"], "M-4")
    _add_body(document, "当 d_i=0 时令后验上限为 1。若候选降低目标概率，最终下降恰为 ρ̂_i d_i；ρ̂_i 不超过 ε_i/d_i，因此逐图后验下降不超过预算。该软概率界不等价于总体目标召回率保证。")
    _equation(document, formulas["margin"], "M-5")
    _add_body(document, "锚点 A_i=1 要求 q0 的 argmax 为目标、q0,T≥τp 且目标与最大非目标的初始间隔至少为 τm。式 M-5 仅作用于锚点；不存在正 Δ 的类别时取 1，非锚点也取 1，且 0<δ≤τm。每个非目标类别的融合间隔为初始间隔减去 ρ̂_i Δ_i,k，故上限可逐类解析求解。")
    _equation(document, formulas["fusion"], "M-6")
    _equation(document, formulas["projection"], "M-7")
    _add_body(document, "命题 M.1（可行投影）：上述线性不等式与 0≤ρ≤1 的交集为闭区间 [0,min(cpost,cmargin)]，必含回退点 0。因此式 M-6 是原始路由在该区间上的唯一欧氏投影，也是不超过原始路由的最大可行采用幅度。命题 M.2（锚点保持）：δ>0 保证锚点目标 argmax 不变。证明只依赖后验与区间约束，不依赖骨干架构；它同样保持假阳性锚点，必须披露这一代价。凸组合还保证概率单纯形有效及 NLL 的凸性上界，但不保证超过最优端点专家。")

    _add_heading(document, "M.4 非冗余条件与专家互补性上界", 2)
    _equation(document, formulas["dominance"], "M-8")
    _add_body(document, "推论 M.3（充分冗余条件）：非目标最大概率不超过 1-qT，因而式 M-8 成立。如果右侧阈值条件成立，后验约束已经保证所有目标锚点间隔严格超过 δ，间隔层不再改变最终路由。初轮 τp≥0.60、εmax≤0.08、δ≤0.01 均落入该区间，因此不能以该轮锚点保持率证明第二层有价值。违反此充分条件只是允许非冗余，不保证真实样本中会激活。")
    _equation(document, formulas["capacity"], "M-9")
    _equation(document, formulas["recall"], "M-10")
    _add_body(document, "命题 M.4（固定双专家分类容量）：假设两端后验都具有唯一最大类别。两端最大类别相同的样本，对任意凸路由仍保持该类别，因为其逐类正间隔的凸组合仍为正。分类变化因此只能来自 D_i=1 的样本，得到式 M-9；限制到真实目标子集得到式 M-10。它是给定冻结专家与评价样本的容量上界，不是泛化误差定理。它也不限制概率 NLL 改善，因为同一预测类别内部仍可改变真实类概率。")

    _add_heading(document, "M.5 两轮开发实验与确定性审计", 2)
    _add_body(document, "每轮训练 9 个预算/锚点组合，每个检查点评估 3 个偏移，共 27 个候选。只复用 Fold 0 的冻结专家与候选门；策略拟合、停止和开发评价分别为 340、340、1,696 张。初轮仅 14 张图片触发后验上限，间隔独立激活为 0；其旧锁已废止，不能启动确证。第二轮降低锚点阈值以离开充分冗余区域，但 27 个候选仍全部未通过独立间隔作用判据。")
    rows = []
    for name, label in (("A0", "回退 q0"), ("A1", "候选 qφ"), ("A3", "历史固定预算"), ("A4", "仅后验约束"), ("A6", "完整 ABMP（诊断）")):
        m = r2["methods"][name]["metrics"]
        rows.append([label, f"{m['nll']:.6f}", _percent(m["macro_f1"]), _percent(m["target_recall"]), _percent(m["ti_intrusion_to_target"]), _percent(m["metallic_intrusion_to_target"])])
    _table(document, ["方法", "NLL", "Macro F1", "目标召回", "含钛误入", "金属误入"], rows, [4.0, 2.0, 2.5, 2.5, 2.5, 2.5])
    _add_body(document, "表中 A6 是开发失败后的最低 NLL 诊断配置（εmax=0.08、τp=0.40、τm=0.05、δ=0.005、λ=-0.5），不是晋级选中模型；A4 与 A6 的逐图输出完全相同。其 NLL 相对 q0 下降约 0.002780 nat，但目标召回下降 0.904 个百分点，含钛误入不变，金属误入下降 0.847 个百分点。开发点估计不承担显著性或非劣效确证。同均值固定预算重放使用同一路由权重，仅替换预算，不是重新训练的固定预算对照。")
    a = r2["methods"]["A6"]["audit"]
    _table(document, ["检查项", "观察值", "证据含义"], [
        ["真实数据不变量", "全部违规计数为 0", "约束实现符合定义"],
        ["目标锚点", f"{a['anchor_count']}；真目标 {a['anchor_true_target_count']}；假阳性 {a['anchor_false_positive_count']}", "保持率 100%，同时固定误入"],
        ["预算分布", f"均值 {a['epsilon_mean']:.6f}；标准差 {a['epsilon_std']:.6f}", "非恒定不等于有贡献"],
        ["同均值固定预算重放", f"最大概率差 {geometry['fixed_mean_budget_max_probability_difference']:.1f}；预测分歧 0", "该诊断配置下预算头未改变输出"],
        ["专家分类一致率", _percent(geometry["q0_candidate_argmax_agreement"]), "只有少量样本可改变类别"],
        ["最大预算/最大路由容量", "独立间隔最多 1 张，且为假阳性", "当前间隔设计缺乏真实目标收益空间"],
        ["随机公式检查", f"{invariants['random_samples']:,} 组；违规为 0", "合成数值检查，不是识别性能"],
    ], [4.6, 5.4, 6.0])

    _add_heading(document, "M.6 失败原因、理论边界与下一轮研究", 2)
    _add_body(document, f"容量诊断发现两端分类分歧仅 {geometry['expert_argmax_disagreement_count']} 张。其中真实目标分歧 {geometry['true_target_disagreement_count']} 张，回退原本未命中的真实目标候选改善空间仅 {geometry['true_target_recall_gain_capacity_count']} 张；对应目标召回增益上界为 {_percent(geometry['target_recall_gain_bound'])}。即使原始路由设为 1、预算设为网格最大值，独立间隔保护也最多作用于一个回退假阳性，而不是恢复真实目标。该结果把下一轮工作由阈值搜索转向专家互补性及锚点可靠性。")
    _add_body(document, "下一轮先在拟合/停止子集上比较可达到的目标召回、误入及概率风险，再提出增强候选专家而非单纯增加门控复杂度的最小方案。候选改进应增加真实目标与含钛/金属干扰之间的有效证据分歧，并用冻结专家对照区分收益来源。预算监督必须加入同平均预算的独立训练对照以及独立间隔消融；锚点保持必须同时报告对真目标和假阳性的影响。若能力审计仍无空间，停止该保护层，不以调低判据制造通过。")
    _add_body(document, "理论深化的可实施主线是：固定专家凸融合的容量约束、约束可行域及非冗余条件、锚点可靠性与风险取舍，再由独立数据验证预算分配的经验作用。以上是可推导、可核验的结构结果，不宣称通用投影、凸性或多任务学习为首次；算法创新仍需与相关文献比较并获得独立确证。阶段条件化决策图、真实送检代价和工业回收率继续列为后续方向，不向公开标本图像附加未经确认的流程标签。")
    _add_heading(document, "M.7 复现证据与晋级边界", 2)
    _add_body(document, "代码、配置、逐图预测、训练历史、数据及权重哈希、几何诊断与结构图源进入版本控制；原始图片和模型权重不公开分发。重新生成路径为 scripts/run_abmp_rsg_development.py、scripts/analyze_abmp_development.py、scripts/generate_abmp_rsg_figure.py 与 tools/append_abmp_to_official_report.py。旧初轮锁通过 protocol_status.json 标记为废止；第二轮不生成锁，Fold 1/2 不启动。新增数学性质与负结果可充实技术报告，但目前不能将 ABMP 的总体性能优势列为已完成创新成果。")
    _mark_owned_appendix(document, title_node)
    temporary = report.with_suffix(".abmp.tmp.docx")
    document.save(temporary)
    temporary.replace(report)
    return True


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", type=Path, default=REPORT)
    parser.add_argument("--refresh", action="store_true", help="Rebuild only the final appendix M; refuse later appendices.")
    args = parser.parse_args()
    source = ROOT / "outputs/training/abmp_rsg_v2"
    summaries = [json.loads((source / name / "development_summary.json").read_text(encoding="utf-8")) for name in ("development_fold_0", "development_fold_0_r2")]
    geometry = json.loads((ROOT / "outputs/theory/abmp_rsg_development_geometry.json").read_text(encoding="utf-8"))
    invariants = json.loads((ROOT / "outputs/theory/abmp_rsg_invariants.json").read_text(encoding="utf-8"))
    print(json.dumps({"changed": append_appendix(args.report, summaries, geometry, invariants, ROOT / "outputs/paper_figures_v5", refresh=args.refresh)}, indent=2))


if __name__ == "__main__":
    main()
