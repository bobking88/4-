import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))

from append_oos_rsg_to_official_report import TITLE, append_appendix


class OOSRSGReportAppendixTests(unittest.TestCase):
    def test_appends_idempotent_evidence_backed_appendix(self):
        root = Path(__file__).resolve().parents[1]
        analysis = json.loads(
            (root / "outputs/training/seen_unseen_gate_study_v1/analysis.json").read_text(encoding="utf-8")
        )
        degeneracy = json.loads(
            (root / "outputs/theory/oos_gate_degeneracy.json").read_text(encoding="utf-8")
        )
        figure = root / "outputs/paper_figures_v3/fig_oos_rsg_training_architecture.png"

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.docx"
            document = Document()
            document.add_heading("Existing report", level=1)
            document.save(path)

            changed = append_appendix(path, figure, analysis, degeneracy)
            self.assertTrue(changed)
            first = Document(path)
            self.assertEqual(sum(p.text == TITLE for p in first.paragraphs), 1)
            self.assertEqual(len(first.inline_shapes), 1)
            self.assertEqual(len(first.tables), 2)
            text = "\n".join(p.text for p in first.paragraphs)
            self.assertIn("条件无偏", text)
            self.assertIn("门控退化", text)
            self.assertIn("不是三个独立专家", text)

            changed_again = append_appendix(path, figure, analysis, degeneracy)
            self.assertFalse(changed_again)
            second = Document(path)
            self.assertEqual(sum(p.text == TITLE for p in second.paragraphs), 1)


if __name__ == "__main__":
    unittest.main()
