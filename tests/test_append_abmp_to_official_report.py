import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))


class ReportTests(unittest.TestCase):
    def _append_fixture(self, report):
        from append_abmp_to_official_report import append_appendix
        summaries = [json.loads((ROOT / "outputs/training/abmp_rsg_v2" / name / "development_summary.json").read_text(encoding="utf-8")) for name in ("development_fold_0", "development_fold_0_r2")]
        geometry = json.loads((ROOT / "outputs/theory/abmp_rsg_development_geometry.json").read_text(encoding="utf-8"))
        invariants = json.loads((ROOT / "outputs/theory/abmp_rsg_invariants.json").read_text(encoding="utf-8"))
        return lambda **kwargs: append_appendix(report, summaries, geometry, invariants, ROOT / "outputs/paper_figures_v5", **kwargs)

    def test_refresh_refuses_later_content_without_heading_style(self):
        with tempfile.TemporaryDirectory() as temporary:
            report = Path(temporary) / "report.docx"
            Document().save(report)
            append = self._append_fixture(report)
            append()
            document = Document(report)
            document.add_paragraph("User appendix N", style="Title")
            document.add_paragraph("User results must remain.")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "User table"
            document.save(report)
            before = report.read_bytes()
            with self.assertRaises(ValueError):
                append(refresh=True)
            self.assertEqual(report.read_bytes(), before)

    def test_refresh_preserves_old_paragraph_with_ordinary_line_break(self):
        from append_abmp_to_official_report import TITLE
        with tempfile.TemporaryDirectory() as temporary:
            report = Path(temporary) / "report.docx"
            document = Document()
            paragraph = document.add_paragraph("Old content stays.")
            paragraph.add_run().add_break()
            document.save(report)
            append = self._append_fixture(report)
            append()
            document = Document(report)
            title = next(p for p in document.paragraphs if p.text == TITLE)
            page_break = title._p.getprevious()
            page_break.getparent().remove(page_break)
            document.save(report)
            append(refresh=True)
            self.assertIn("Old content stays.", "\n".join(p.text for p in Document(report).paragraphs))

    def test_refresh_refuses_user_edits_inside_generated_appendix(self):
        with tempfile.TemporaryDirectory() as temporary:
            report = Path(temporary) / "report.docx"
            Document().save(report)
            append = self._append_fixture(report)
            append()
            document = Document(report)
            document.paragraphs[-1].add_run(" Manual conclusion.")
            document.save(report)
            before = report.read_bytes()
            with self.assertRaises(ValueError):
                append(refresh=True)
            self.assertEqual(report.read_bytes(), before)

    def test_refresh_refuses_changed_embedded_image_with_same_relationship(self):
        from io import BytesIO
        from PIL import Image
        with tempfile.TemporaryDirectory() as temporary:
            report = Path(temporary) / "report.docx"
            Document().save(report)
            append = self._append_fixture(report)
            append()
            document = Document(report)
            embedded = document.inline_shapes[0]._inline.xpath(".//a:blip")[0]
            relationship = embedded.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            buffer = BytesIO()
            Image.new("RGB", (3, 3), "black").save(buffer, format="PNG")
            document.part.related_parts[relationship]._blob = buffer.getvalue()
            document.save(report)
            before = report.read_bytes()
            with self.assertRaises(ValueError):
                append(refresh=True)
            self.assertEqual(report.read_bytes(), before)

    def test_preserves_report_and_appends_failed_development_idempotently(self):
        from append_abmp_to_official_report import TITLE, append_appendix

        summaries = [json.loads((ROOT / "outputs/training/abmp_rsg_v2" / name / "development_summary.json").read_text(encoding="utf-8")) for name in ("development_fold_0", "development_fold_0_r2")]
        geometry = json.loads((ROOT / "outputs/theory/abmp_rsg_development_geometry.json").read_text(encoding="utf-8"))
        invariants = json.loads((ROOT / "outputs/theory/abmp_rsg_invariants.json").read_text(encoding="utf-8"))
        with tempfile.TemporaryDirectory() as temporary:
            report = Path(temporary) / "report.docx"
            document = Document()
            document.add_heading("Existing report", 1)
            document.add_paragraph("Preserve existing results.")
            document.save(report)
            self.assertTrue(append_appendix(report, summaries, geometry, invariants, ROOT / "outputs/paper_figures_v5"))
            updated = Document(report)
            text = "\n".join(p.text for p in updated.paragraphs)
            for phrase in ("Preserve existing results.", "未通过", "Fold 1/2", "不是独立确证", "假阳性", "互补性", "固定预算", "2,418"):
                self.assertIn(phrase, text)
            self.assertEqual(sum(p.text == TITLE for p in updated.paragraphs), 1)
            self.assertGreaterEqual(len(updated.inline_shapes), 8)
            self.assertEqual(len(updated.tables), 3)
            self.assertFalse(append_appendix(report, summaries, geometry, invariants, ROOT / "outputs/paper_figures_v5"))
            self.assertTrue(append_appendix(report, summaries, geometry, invariants, ROOT / "outputs/paper_figures_v5", refresh=True))
            refreshed = Document(report)
            self.assertEqual(sum(p.text == TITLE for p in refreshed.paragraphs), 1)
            self.assertEqual(len(refreshed.inline_shapes), len(updated.inline_shapes))
            refreshed.add_heading("User appendix N", 1)
            refreshed.save(report)
            with self.assertRaises(ValueError):
                append_appendix(report, summaries, geometry, invariants, ROOT / "outputs/paper_figures_v5", refresh=True)

    def test_refuses_mislabeled_confirmation(self):
        from append_abmp_to_official_report import append_appendix
        with self.assertRaises(ValueError):
            append_appendix(Path("not_read.docx"), [{"purpose": "confirmation"}], {}, {}, Path("missing"))


if __name__ == "__main__":
    unittest.main()
