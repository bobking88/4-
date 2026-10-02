import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from append_abmp_capacity_to_report import TITLE, append_capacity_appendix


class CapacityReportTests(unittest.TestCase):
    def setUp(self):
        self.summary = json.loads((ROOT / "outputs/theory/abmp_candidate_capacity_v1/audit_summary.json").read_text(encoding="utf-8"))
        self.properties = json.loads((ROOT / "outputs/theory/abmp_candidate_capacity_v1/geometry_properties.json").read_text(encoding="utf-8"))

    def assets(self, folder):
        (folder / "capacity_formulas").mkdir()
        for name in ("scale", "commutation", "monotonicity", "interval", "capacity", "protected_scale"):
            Image.new("RGB", (400, 40), "white").save(folder / "capacity_formulas" / f"capacity_{name}.png")
        Image.new("RGB", (1000, 400), "white").save(folder / "fig_abmp_candidate_capacity.png")

    def test_append_preserves_all_existing_body_nodes_and_is_idempotent(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            path = folder / "report.docx"
            document = Document()
            document.add_paragraph("Existing report text")
            document.add_table(rows=1, cols=2).cell(0, 0).text = "Existing results"
            document.save(path)
            before = [etree.tostring(node, method="c14n") for node in Document(path)._element.body if node.tag != qn("w:sectPr")]
            self.assets(folder)
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            self.assertTrue(append_capacity_appendix(path, self.summary, self.properties, folder, expected_sha256=digest))
            edited = Document(path)
            after = [etree.tostring(node, method="c14n") for node in edited._element.body if node.tag != qn("w:sectPr")]
            self.assertEqual(before, after[:len(before)])
            self.assertEqual(sum(p.text == TITLE for p in edited.paragraphs), 1)
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            self.assertFalse(append_capacity_appendix(path, self.summary, self.properties, folder, expected_sha256=digest))
            self.assertEqual(digest, hashlib.sha256(path.read_bytes()).hexdigest())

    def test_wrong_expected_hash_refuses_edit(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.docx"
            Document().save(path)
            with self.assertRaisesRegex(ValueError, "hash"):
                append_capacity_appendix(path, self.summary, self.properties, Path(directory), expected_sha256="0"*64)

    def test_outer_data_cannot_be_presented_as_development_evidence(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.docx"
            Document().save(path)
            self.summary["outer_images_loaded"] = True
            with self.assertRaisesRegex(ValueError, "development"):
                append_capacity_appendix(path, self.summary, self.properties, Path(directory), expected_sha256=hashlib.sha256(path.read_bytes()).hexdigest())

    def test_changed_observation_cannot_reuse_fixed_diagnostic_narrative(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            path = folder / "report.docx"
            Document().save(path)
            self.assets(folder)
            self.summary["subsets"]["projector_stop"]["segment_capacity"]["full_verified"]["true_target_feasible"] = 46
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            with self.assertRaisesRegex(ValueError, "observations"):
                append_capacity_appendix(path, self.summary, self.properties, folder, expected_sha256=digest)
            self.assertEqual(digest, hashlib.sha256(path.read_bytes()).hexdigest())


if __name__ == "__main__":
    unittest.main()
