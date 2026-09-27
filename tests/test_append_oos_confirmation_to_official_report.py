from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "tools"))


class OOSConfirmationReportAppendixTests(unittest.TestCase):
    def test_appends_confirmation_and_constraint_appendix_idempotently(self) -> None:
        from append_oos_confirmation_to_official_report import TITLE, append_appendix

        analysis = json.loads(
            (
                PROJECT_ROOT
                / "outputs"
                / "training"
                / "oos_rsg_confirmation_v1"
                / "analysis"
                / "analysis.json"
            ).read_text(encoding="utf-8")
        )
        audit = json.loads(
            (
                PROJECT_ROOT
                / "outputs"
                / "training"
                / "oos_rsg_confirmation_manifests_v1"
                / "audit.json"
            ).read_text(encoding="utf-8")
        )
        figure = (
            PROJECT_ROOT
            / "outputs"
            / "paper_figures_v3"
            / "fig_oos_rsg_confirmation.png"
        )

        with tempfile.TemporaryDirectory() as temporary:
            report = Path(temporary) / "report.docx"
            document = Document()
            document.add_heading("Existing report", level=1)
            document.save(report)

            changed = append_appendix(report, figure, analysis, audit)
            self.assertTrue(changed)

            updated = Document(report)
            self.assertEqual(sum(p.text == TITLE for p in updated.paragraphs), 1)
            self.assertEqual(len(updated.inline_shapes), 1)
            self.assertEqual(len(updated.tables), 3)
            text = "\n".join(p.text for p in updated.paragraphs)
            for phrase in (
                "2.476 个百分点",
                "不晋级",
                "TC-OOS-RSG",
                "目标软漏识风险",
                "现有 final_eval 已经被读取",
                "下表中效应统一定义为",
            ):
                self.assertIn(phrase, text)
            self.assertNotIn("表 L-2 中效应统一定义为", text)

            changed_again = append_appendix(report, figure, analysis, audit)
            self.assertFalse(changed_again)
            rerun = Document(report)
            self.assertEqual(sum(p.text == TITLE for p in rerun.paragraphs), 1)


if __name__ == "__main__":
    unittest.main()
